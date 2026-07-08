"""GRN Dispatch Note Excel writer.

Recreates the exact layout, styling and pagination of the resort's official
"GRN DISPATCH NOTE" workbook (one sheet per month, 30 line-items per printed
dispatch page, Outrigger logo top-right of every page, signature block,
landscape print layout with a manual page-break every 44 rows).

Public entry point: write_grn_dispatch_workbook(...)
"""
from __future__ import annotations

import os
import re
from datetime import datetime
from typing import Dict, Iterable, List, Optional, Tuple

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.pagebreak import Break

# ---------------------------------------------------------------------------
# Layout constants -- measured directly from the reference workbook
# "GRN DISPATCH NOTE APRIL - 2026.xlsx" so the output is pixel/cell-for-cell
# the same template.
# ---------------------------------------------------------------------------
ROWS_PER_PAGE = 30          # line items per dispatch page
BLOCK_HEIGHT = 44           # total rows (incl. blanks + signature) per page
CURRENCIES = ["usd", "mvr", "eur", "gbp", "sgd"]
CURRENCY_LABELS = ["USD", "MVR", "EUR", "GBP", "SGD"]
CURRENCY_COLS = [6, 7, 8, 9, 10]   # F, G, H, I, J (1-indexed)

ACCOUNTING_FMT = '_-* #,##0.00_-;\\-* #,##0.00_-;_-* "-"??_-;_-@_-'
DATE_FMT = "d-mmm-yy"

HEADER_FILL = PatternFill(fill_type="solid", fgColor="E7E6E6")
WHITE_FILL = PatternFill(fill_type="solid", fgColor="FFFFFF")

FONT_BASE = "Arial Narrow"
COL_WIDTHS = {"A": 4.5546875, "B": 9.33203125, "C": 37.44140625,
              "D": 30.44140625, "E": 18.88671875, "F": 9.88671875,
              "G": 11.109375, "H": 9.109375, "K": 45.0}

LOGO_ANCHOR_COL = "K"
LOGO_SIZE = (468, 91)  # px, native size of the reference logo image

PLACEHOLDER_PO = "MAM-0000"
PLACEHOLDER_GRN = "RC-MAM-0000"

_MEDIUM = Side(style="medium")
_THIN = Side(style="thin")
_DOTTED = Side(style="dotted")


def _month_year_from_date_str(value) -> Tuple[int, int]:
    """Best-effort (year, month) extraction from a row's date value.
    Accepts 'DD.MM.YYYY' strings (the app's native format), datetime/date
    objects, or falls back to today's month/year if unparseable."""
    if isinstance(value, datetime):
        return value.year, value.month
    if hasattr(value, "year") and hasattr(value, "month"):
        return value.year, value.month
    if isinstance(value, str):
        s = value.strip()
        m = re.match(r"^(\d{1,2})[.\-/](\d{1,2})[.\-/](\d{4})$", s)
        if m:
            day, month, year = int(m.group(1)), int(m.group(2)), int(m.group(3))
            if 1 <= month <= 12:
                return year, month
    now = datetime.now()
    return now.year, now.month


def _group_rows_by_month(rows: List[Dict]) -> "dict":
    """Groups rows into one bucket per (year, month) found in their date
    field -- this is what lets the export intelligently split a mixed batch
    of invoices into one correctly-named sheet per month, exactly like the
    reference workbook (a fresh sheet per month, e.g. 'APRIL')."""
    buckets: Dict[Tuple[int, int], List[Dict]] = {}
    for r in rows:
        key = _month_year_from_date_str(r.get("date"))
        buckets.setdefault(key, []).append(r)
    # Sort months chronologically, and rows within a month by date then by
    # original order (stable sort keeps AI-extraction order for same date).
    for key in buckets:
        buckets[key].sort(key=lambda r: _date_sort_key(r.get("date")))
    return dict(sorted(buckets.items(), key=lambda kv: kv[0]))


def _date_sort_key(value):
    y, m = 0, 0
    d = 0
    if isinstance(value, str):
        mm = re.match(r"^(\d{1,2})[.\-/](\d{1,2})[.\-/](\d{4})$", value.strip())
        if mm:
            d, m, y = int(mm.group(1)), int(mm.group(2)), int(mm.group(3))
    elif hasattr(value, "year"):
        y, m, d = value.year, value.month, getattr(value, "day", 0)
    return (y, m, d)


def _chunk(seq: List[Dict], size: int) -> Iterable[List[Dict]]:
    for i in range(0, len(seq), size):
        yield seq[i:i + size]


def _sheet_name(year: int, month: int) -> str:
    return datetime(year, month, 1).strftime("%B").upper()


def _set_col_widths(ws):
    for col, width in COL_WIDTHS.items():
        ws.column_dimensions[col].width = width


def _font(size=10, bold=False, italic=False):
    return Font(name=FONT_BASE, size=size, bold=bold, italic=italic, color="FF000000")


def _center(wrap=False, vertical="center"):
    return Alignment(horizontal="center", vertical=vertical, wrap_text=wrap)


def _apply_box_border(ws, row, col, left=None, right=None, top=None, bottom=None):
    cell = ws.cell(row=row, column=col)
    cell.border = Border(left=left or Side(style=None), right=right or Side(style=None),
                          top=top or Side(style=None), bottom=bottom or Side(style=None))


def _write_page(ws, block_start_row: int, month_label: str, year: int,
                 dispatch_no: int, batch: List[Dict], logo_path: Optional[str]):
    """Writes ONE 44-row dispatch page starting at `block_start_row`
    (1-indexed), matching the reference template exactly."""
    title_row = block_start_row + 3
    date_row = block_start_row + 5
    hdr1_row = block_start_row + 7
    hdr2_row = hdr1_row + 1
    data_start = hdr1_row + 2
    sig_row = block_start_row + 42
    handed_row = block_start_row + 43

    # -- Title (merged, 2 rows tall, cols A:G) --------------------------
    ws.merge_cells(start_row=title_row, start_column=1, end_row=title_row + 1, end_column=7)
    tcell = ws.cell(row=title_row, column=1, value=f"GRN DISPATCH {month_label}- {year}")
    tcell.font = _font(size=20, bold=True, italic=True)
    tcell.alignment = _center(wrap=True)
    tcell.fill = WHITE_FILL
    tcell.border = Border(left=_MEDIUM, right=_THIN, top=_THIN, bottom=_THIN)
    ws.row_dimensions[title_row].height = 15
    ws.row_dimensions[title_row + 1].height = 15

    # -- DATE / DISPATCH row --------------------------------------------
    ws.row_dimensions[date_row].height = 15
    lbl = ws.cell(row=date_row, column=1, value="DATE")
    lbl.font = _font(bold=True)
    lbl.alignment = _center()
    lbl.fill = WHITE_FILL
    lbl.border = Border(left=_MEDIUM, top=_THIN, bottom=_THIN)
    lbl.number_format = DATE_FMT

    dval = ws.cell(row=date_row, column=3, value="=TODAY()")
    dval.font = _font(bold=True, italic=True)
    dval.alignment = _center()
    dval.fill = WHITE_FILL
    dval.border = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
    dval.number_format = DATE_FMT

    ws.merge_cells(start_row=date_row, start_column=6, end_row=date_row, end_column=7)
    dlbl = ws.cell(row=date_row, column=6, value="DISPATCH")
    dlbl.font = _font(bold=True)
    dlbl.alignment = _center(wrap=True)
    dlbl.fill = WHITE_FILL
    dlbl.border = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
    dlbl.number_format = ACCOUNTING_FMT

    dno = ws.cell(row=date_row, column=8, value=dispatch_no)
    dno.font = _font(bold=True)
    dno.alignment = _center()
    dno.fill = WHITE_FILL
    dno.border = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)

    # -- Header row 1 (merged 2 rows for most cols) ----------------------
    headers = ["NO - #", "INVOICE DATE", "SUPPLIER NAME", "PURCHASE ORDER #", "INVOICE #"]
    for i, text in enumerate(headers, start=1):
        ws.merge_cells(start_row=hdr1_row, start_column=i, end_row=hdr2_row, end_column=i)
        c = ws.cell(row=hdr1_row, column=i, value=text)
        c.font = _font(bold=True)
        c.alignment = _center(wrap=True)
        c.fill = HEADER_FILL
        c.number_format = DATE_FMT
        left = _MEDIUM if i == 1 else _THIN
        c.border = Border(left=left, right=_THIN, top=_THIN, bottom=_THIN)

    ws.merge_cells(start_row=hdr1_row, start_column=6, end_row=hdr1_row, end_column=10)
    gamt = ws.cell(row=hdr1_row, column=6, value="GRN AMOUNT")
    gamt.font = _font(bold=True)
    gamt.alignment = _center(wrap=True)
    gamt.fill = HEADER_FILL
    gamt.number_format = ACCOUNTING_FMT
    gamt.border = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)

    ws.merge_cells(start_row=hdr1_row, start_column=11, end_row=hdr2_row, end_column=11)
    grnno = ws.cell(row=hdr1_row, column=11, value="GRN NO.")
    grnno.font = _font(bold=True)
    grnno.alignment = _center(wrap=True)
    grnno.fill = HEADER_FILL
    grnno.number_format = DATE_FMT
    grnno.border = Border(left=_THIN, right=_MEDIUM, top=_THIN, bottom=_THIN)

    for i, label in zip(CURRENCY_COLS, CURRENCY_LABELS):
        c = ws.cell(row=hdr2_row, column=i, value=label)
        c.font = _font(bold=True)
        c.alignment = _center(wrap=True)
        c.fill = HEADER_FILL
        c.number_format = ACCOUNTING_FMT
        c.border = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)

    # -- Data rows (always exactly ROWS_PER_PAGE, padded with placeholders) --
    for i in range(ROWS_PER_PAGE):
        row = data_start + i
        rec = batch[i] if i < len(batch) else None

        no_cell = ws.cell(row=row, column=1, value=i + 1)
        no_cell.font = _font()
        no_cell.alignment = _center()
        no_cell.fill = WHITE_FILL
        no_cell.border = Border(left=_MEDIUM, right=_DOTTED, top=_DOTTED, bottom=_DOTTED)

        inv_date = rec.get("date", "") if rec else ""
        b = ws.cell(row=row, column=2, value=inv_date)
        b.font = _font()
        b.alignment = _center()
        b.fill = WHITE_FILL
        b.border = Border(left=_DOTTED, right=_DOTTED, top=_DOTTED, bottom=_DOTTED)

        supplier = (rec.get("supplier", "") or "UNKNOWN SUPPLIER") if rec else ""
        c = ws.cell(row=row, column=3, value=supplier)
        c.font = _font()
        c.alignment = Alignment(horizontal="left", vertical=None)
        c.fill = WHITE_FILL
        c.border = Border(left=_DOTTED, right=_DOTTED, top=_DOTTED, bottom=_DOTTED)

        po = (rec.get("po", "") or PLACEHOLDER_PO) if rec else PLACEHOLDER_PO
        d = ws.cell(row=row, column=4, value=po)
        d.font = _font()
        d.alignment = _center()
        d.fill = WHITE_FILL
        d.border = Border(left=_DOTTED, right=_DOTTED, top=_DOTTED, bottom=_DOTTED)

        invoice = rec.get("invoice", "") if rec else ""
        if isinstance(invoice, str) and invoice.upper() == "NO-INVOICE":
            invoice = ""
        e = ws.cell(row=row, column=5, value=invoice)
        e.font = _font()
        e.alignment = Alignment(horizontal="left", vertical=None)
        e.fill = WHITE_FILL
        e.border = Border(left=_DOTTED, right=_DOTTED, top=_DOTTED, bottom=_DOTTED)

        for col_idx, cur in zip(CURRENCY_COLS, CURRENCIES):
            val = rec.get(cur, "") if rec else ""
            cc = ws.cell(row=row, column=col_idx)
            cc.font = _font()
            cc.alignment = _center()
            cc.fill = WHITE_FILL
            cc.border = Border(left=_DOTTED, right=_DOTTED, top=_DOTTED, bottom=_DOTTED)
            cc.number_format = ACCOUNTING_FMT
            if val not in ("", None):
                try:
                    cc.value = float(str(val).replace(",", ""))
                except (ValueError, TypeError):
                    cc.value = val
            else:
                cc.value = ""

        grn = (rec.get("grn", "") or PLACEHOLDER_GRN) if rec else PLACEHOLDER_GRN
        k = ws.cell(row=row, column=11, value=grn)
        k.font = _font()
        k.alignment = _center()
        k.fill = WHITE_FILL
        k.border = Border(left=_DOTTED, right=_MEDIUM, top=_DOTTED, bottom=_DOTTED)

    # -- Signature block --------------------------------------------------
    sdots1 = ws.cell(row=sig_row, column=3, value="\u2026...............................................")
    sdots1.font = _font(size=11, bold=True)
    sdots1.alignment = _center()

    sdots2 = ws.cell(row=sig_row, column=5, value="\u2026..........................................")
    sdots2.font = _font(size=11, bold=True)
    sdots2.alignment = Alignment(horizontal="center")

    handed = ws.cell(row=handed_row, column=3, value="HANDED OVER BY")
    handed.font = _font(size=11, bold=True)
    handed.alignment = Alignment(horizontal="center")
    handed.border = Border(bottom=_MEDIUM)

    received = ws.cell(row=handed_row, column=5, value="RECEIVED BY")
    received.font = _font(size=11, bold=True)
    received.alignment = Alignment(horizontal="center")
    received.border = Border(bottom=_MEDIUM)

    ws.row_dimensions[hdr1_row].height = 15
    ws.row_dimensions[handed_row].height = 15

    # -- Outrigger logo, top-right of this page ---------------------------
    if logo_path and os.path.isfile(logo_path):
        img = XLImage(logo_path)
        img.width, img.height = LOGO_SIZE
        ws.add_image(img, f"{LOGO_ANCHOR_COL}{block_start_row + 1}")


def _write_month_sheet(wb: Workbook, month_label: str, year: int, rows: List[Dict],
                        start_dispatch_no: int, logo_path: Optional[str]) -> int:
    sheet_name = _sheet_name(year, month := datetime.strptime(month_label, "%B").month)[:31]
    ws = wb.create_sheet(title=_sheet_name(year, month))
    ws.sheet_view.showGridLines = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(left=0.5, right=0.5, top=0.75, bottom=0.75, header=0.3, footer=0.3)
    _set_col_widths(ws)
    ws.row_dimensions[1].height = 15

    batches = list(_chunk(rows, ROWS_PER_PAGE)) or [[]]
    dispatch_no = start_dispatch_no
    for page_idx, batch in enumerate(batches):
        block_start_row = 1 + page_idx * BLOCK_HEIGHT
        _write_page(ws, block_start_row, month_label, year, dispatch_no, batch, logo_path)
        dispatch_no += 1
        if page_idx < len(batches) - 1:
            ws.row_breaks.append(Break(id=block_start_row + BLOCK_HEIGHT - 1))

    last_row = len(batches) * BLOCK_HEIGHT
    ws.print_area = f"A1:K{last_row}"
    return dispatch_no


def write_grn_dispatch_workbook(rows: List[Dict], start_dispatch_no: int,
                                 out_path: str, logo_path: Optional[str] = None) -> Tuple[str, int]:
    """Builds the full multi-sheet GRN Dispatch Note workbook -- one sheet
    per month found in `rows` (e.g. 'APRIL', 'MAY'), each laid out exactly
    like the reference "GRN DISPATCH NOTE" workbook: 30 line-items per
    printed page, a running DISPATCH number, the Outrigger logo top-right of
    every page, and a signature block, ready to print in landscape.

    rows: list of dicts with keys date, supplier, po, invoice, usd, mvr,
          eur, gbp, sgd, grn (same shape the AI Extract tab already uses).
    start_dispatch_no: first DISPATCH number to use; increments by 1 for
          every page written (across all months, in chronological order).
    Returns (saved_path, next_dispatch_no) so the caller can persist the
    next starting number for the following export.
    """
    wb = Workbook()
    wb.remove(wb.active)  # drop the default blank sheet

    grouped = _group_rows_by_month(rows)
    dispatch_no = start_dispatch_no
    for (year, month), month_rows in grouped.items():
        month_label = datetime(year, month, 1).strftime("%B").upper()
        dispatch_no = _write_month_sheet(wb, month_label, year, month_rows, dispatch_no, logo_path)

    out_dir = os.path.dirname(out_path)
    if out_dir:
        os.makedirs(out_dir, exist_ok=True)
    base, ext = os.path.splitext(out_path)
    final_path = out_path
    cnt = 1
    while os.path.exists(final_path):
        final_path = f"{base}_{cnt}{ext}"
        cnt += 1
    wb.save(final_path)
    return final_path, dispatch_no
